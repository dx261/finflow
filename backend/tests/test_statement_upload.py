"""Statement upload and import pipeline integration tests."""

import io
import uuid
from decimal import Decimal

from app.core.config import settings
from app.core.database import AsyncSessionLocal
from app.modules.accounts.models import Account
from app.modules.relations.models import RelationType, TransactionRelation
from app.modules.statements.models import (
    StatementImport,
    StatementImportStatus,
    Transaction,
    TransactionType,
)
from app.modules.statements.service import StatementService
from app.modules.transfers.service import TransferService
from openpyxl import Workbook
from sqlalchemy import func, select


def _icbc_xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["中国工商银行", None, None, None, None])
    ws.append(["账号", "6222020", None, None, None])
    ws.append(["交易日期", "摘要", "借方金额", "贷方金额", "余额"])
    ws.append(["2026-09-01", "工资", None, 8000.0, 50000.0])
    ws.append(["2026-09-02", "购物", 120.5, None, 49879.5])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _single_transaction_xlsx_bytes(
    account_number: str, *, debit: float | None = None, credit: float | None = None
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["中国工商银行", None, None, None, None, None])
    ws.append(["账号", account_number, None, None, None, None])
    ws.append(["交易日期", "摘要", "交易对手", "借方金额", "贷方金额", "余额"])
    ws.append(["2026-09-05", "内部转账", "本人", debit, credit, 1000.0])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _balance_rows_xlsx_bytes(rows: list[tuple[str, float, float]]) -> bytes:
    """Build an ICBC-like sheet from (date, debit, balance) rows."""
    wb = Workbook()
    ws = wb.active
    ws.append(["中国工商银行", None, None, None, None])
    ws.append(["账号", "6222020", None, None, None])
    ws.append(["交易日期", "摘要", "借方金额", "贷方金额", "余额"])
    for txn_date, debit, balance in rows:
        ws.append([txn_date, f"消费-{txn_date}", debit, None, balance])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def _create_account(client, auth_headers) -> str:
    banks = (await client.get("/api/v1/banks", headers=auth_headers)).json()
    bank_id = next(bank["id"] for bank in banks if bank["code"] == "ICBC")
    response = await client.post(
        "/api/v1/accounts",
        headers=auth_headers,
        json={
            "bank_id": bank_id,
            "name": "工行储蓄卡",
            "account_type": "CHECKING",
            "currency": "CNY",
            "actual_balance": "0",
            "minimum_balance": "0",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


async def _upload(client, auth_headers, account_id, content, force=False) -> dict:
    params = {"force_import": "true"} if force else {}
    response = await client.post(
        "/api/v1/statements/upload",
        params=params,
        headers=auth_headers,
        data={"account_id": account_id},
        files={
            "file": (
                "icbc_statement.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    return response


async def _process(import_id: str) -> None:
    async with AsyncSessionLocal() as session:
        await StatementService(session).process_import(uuid.UUID(import_id))


async def test_upload_requires_auth(client, db_reset) -> None:
    response = await _upload(client, {}, str(123), b"x")
    assert response.status_code == 401


async def test_upload_returns_202_and_duplicate_conflicts(
    client, seeded_banks, seeded_categories, auth_headers
) -> None:
    account_id = await _create_account(client, auth_headers)
    content = _icbc_xlsx_bytes()

    first = await _upload(client, auth_headers, account_id, content)
    assert first.status_code == 202, first.text
    body = first.json()
    assert body["is_duplicate"] is False
    assert body["status"] == "UPLOADED"

    # Same file, same account -> duplicate.
    second = await _upload(client, auth_headers, account_id, content)
    assert second.status_code == 409
    assert second.json()["code"] == "DUPLICATE_IMPORT"

    # force_import bypasses the duplicate check.
    forced = await _upload(client, auth_headers, account_id, content, force=True)
    assert forced.status_code == 202
    assert forced.json()["is_duplicate"] is False
    assert forced.json()["import_id"] != body["import_id"]


async def test_upload_rejects_unsupported_extension(client, seeded_banks, auth_headers) -> None:
    account_id = await _create_account(client, auth_headers)
    response = await client.post(
        "/api/v1/statements/upload",
        headers=auth_headers,
        data={"account_id": account_id},
        files={"file": ("notes.txt", b"not a statement", "text/plain")},
    )
    assert response.status_code == 400


async def test_upload_enforces_backend_size_limit(
    client, seeded_banks, auth_headers, monkeypatch
) -> None:
    account_id = await _create_account(client, auth_headers)
    monkeypatch.setattr(settings, "max_upload_size_mb", 0)
    response = await _upload(client, auth_headers, account_id, b"x")
    assert response.status_code == 413
    assert response.json()["code"] == "PAYLOAD_TOO_LARGE"


async def test_process_import_inserts_transactions_and_updates_balance(
    client, seeded_banks, seeded_categories, auth_headers
) -> None:
    account_id = await _create_account(client, auth_headers)
    upload = await _upload(client, auth_headers, account_id, _icbc_xlsx_bytes())
    import_id = upload.json()["import_id"]

    await _process(import_id)

    async with AsyncSessionLocal() as session:
        record = await session.get(StatementImport, uuid.UUID(import_id))
        assert record.status == StatementImportStatus.SUCCESS
        assert record.success_rows == 2
        assert record.failed_rows == 0

        count = await session.scalar(select(func.count()).select_from(Transaction))
        assert count == 2

        account = await session.get(Account, uuid.UUID(account_id))
        assert account.actual_balance == Decimal("49879.5")
        assert account.balance_as_of is not None


async def test_descending_statement_uses_newest_dated_balance(
    client, seeded_banks, seeded_categories, auth_headers
) -> None:
    account_id = await _create_account(client, auth_headers)
    content = _balance_rows_xlsx_bytes(
        [("2026-09-02", 120.5, 49879.5), ("2026-09-01", 10.0, 50000.0)]
    )
    upload = await _upload(client, auth_headers, account_id, content)
    await _process(upload.json()["import_id"])

    async with AsyncSessionLocal() as session:
        account = await session.get(Account, uuid.UUID(account_id))
        assert account.actual_balance == Decimal("49879.5")
        assert account.balance_as_of.date().isoformat() == "2026-09-02"


async def test_older_statement_does_not_overwrite_newer_balance(
    client, seeded_banks, seeded_categories, auth_headers
) -> None:
    account_id = await _create_account(client, auth_headers)
    newer = await _upload(
        client,
        auth_headers,
        account_id,
        _balance_rows_xlsx_bytes([("2026-09-10", 50.0, 1000.0)]),
    )
    await _process(newer.json()["import_id"])
    older = await _upload(
        client,
        auth_headers,
        account_id,
        _balance_rows_xlsx_bytes([("2026-09-05", 20.0, 2000.0)]),
    )
    await _process(older.json()["import_id"])

    async with AsyncSessionLocal() as session:
        account = await session.get(Account, uuid.UUID(account_id))
        assert account.actual_balance == Decimal("1000.0")
        assert account.balance_as_of.date().isoformat() == "2026-09-10"


async def test_process_import_skips_duplicates(
    client, seeded_banks, seeded_categories, auth_headers
) -> None:
    account_id = await _create_account(client, auth_headers)
    upload = await _upload(client, auth_headers, account_id, _icbc_xlsx_bytes())
    import_id = upload.json()["import_id"]

    await _process(import_id)
    await _process(import_id)  # terminal imports are idempotent

    async with AsyncSessionLocal() as session:
        record = await session.get(StatementImport, uuid.UUID(import_id))
        assert record.status == StatementImportStatus.SUCCESS
        assert record.success_rows == 2
        assert record.failed_rows == 0

        count = await session.scalar(select(func.count()).select_from(Transaction))
        assert count == 2  # still only two rows total


async def test_classify_during_import(
    client, seeded_banks, seeded_categories, auth_headers
) -> None:
    account_id = await _create_account(client, auth_headers)
    upload = await _upload(client, auth_headers, account_id, _icbc_xlsx_bytes())
    await _process(upload.json()["import_id"])

    async with AsyncSessionLocal() as session:
        result = await session.scalars(select(Transaction))
        transactions = list(result)
        salary = next(t for t in transactions if t.description == "工资")
        assert salary.category_id is not None


async def test_import_pipeline_persists_internal_transfer_relation(
    client, seeded_banks, seeded_categories, auth_headers
) -> None:
    debit_account_id = await _create_account(client, auth_headers)
    credit_account_id = await _create_account(client, auth_headers)
    debit_upload = await _upload(
        client,
        auth_headers,
        debit_account_id,
        _single_transaction_xlsx_bytes("6222020001", debit=500.0),
    )
    credit_upload = await _upload(
        client,
        auth_headers,
        credit_account_id,
        _single_transaction_xlsx_bytes("6222020002", credit=500.0),
    )

    await _process(debit_upload.json()["import_id"])
    await _process(credit_upload.json()["import_id"])

    async with AsyncSessionLocal() as session:
        # Re-running detection is idempotent and does not reuse either side.
        owner_id = (await session.get(Account, uuid.UUID(debit_account_id))).user_id
        assert await TransferService(session).apply_for_user(owner_id) == []
        relation = await session.scalar(
            select(TransactionRelation).where(
                TransactionRelation.relation_type == RelationType.INTERNAL_TRANSFER
            )
        )
        assert relation is not None
        transactions = list(await session.scalars(select(Transaction)))
        assert len(transactions) == 2
        assert all(t.transaction_type == TransactionType.TRANSFER for t in transactions)
